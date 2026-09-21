import io
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.core.database import get_database_connection
from psycopg2.extras import RealDictCursor

class ReportsExportService:
    def _create_header(self, ws, headers):
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="4F4F4F")
        border = Border(
            left=Side(style='thin', color="DDDDDD"), 
            right=Side(style='thin', color="DDDDDD"), 
            top=Side(style='thin', color="DDDDDD"), 
            bottom=Side(style='thin', color="DDDDDD")
        )
        
        ws.append(headers)
        for cell in ws[ws.max_row]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

    def _auto_adjust_width(self, ws):
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter 
            for cell in col:
                try: 
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = (max_length + 2)
            if adjusted_width > 50:
                adjusted_width = 50
            ws.column_dimensions[column].width = adjusted_width

    def get_all_reports_excel(self):
        wb = openpyxl.Workbook()
        # Remove default sheet
        if 'Sheet' in wb.sheetnames:
            std = wb['Sheet']
            wb.remove(std)
            
        conn = get_database_connection()
        try:
            with conn.cursor(cursor_factory=RealDictCursor) as cursor:
                # --- Pre-fetch general stats for Dashboard/Info ---
                cursor.execute('SELECT COUNT(*) as c FROM job_descriptions')
                total_jobs = cursor.fetchone()['c']
                cursor.execute('SELECT COUNT(*) as c FROM resumes')
                total_resumes = cursor.fetchone()['c']
                cursor.execute("SELECT COUNT(*) as c FROM resumes WHERE processing_status = 'Completed'")
                ai_processed = cursor.fetchone()['c']
                cursor.execute("SELECT COALESCE(AVG(overall_score), 0) as a FROM candidate_scores")
                avg_score = int(cursor.fetchone()['a'])

                # 01_Report_Info
                ws1 = wb.create_sheet("01_Report_Info")
                self._create_header(ws1, ['Field', 'Value'])
                info_data = [
                    ['System', 'ResumeIQ'],
                    ['Report Name', 'Complete Analytics & Reports Export'],
                    ['Generated Date', datetime.now().strftime('%Y-%m-%d')],
                    ['Generated Time', datetime.now().strftime('%H:%M:%S')],
                    ['Total Jobs', total_jobs],
                    ['Total Candidates', total_resumes],
                    ['Total AI Screenings', ai_processed],
                    ['Average AI Match Score', f"{avg_score}%"]
                ]
                for row in info_data:
                    ws1.append(row)
                self._auto_adjust_width(ws1)

                # 02_Dashboard_Summary
                ws2 = wb.create_sheet("02_Dashboard_Summary")
                self._create_header(ws2, ['Metric', 'Value'])
                kpis = [
                    ['Total Jobs', total_jobs],
                    ['Total Candidates', total_resumes],
                    ['Total AI Screenings', ai_processed],
                    ['Average AI Match Score', f"{avg_score}%"],
                    ['', '']
                ]
                for row in kpis:
                    ws2.append(row)
                    
                ws2.append(['--- AI Recommendation Distribution ---', ''])
                self._create_header(ws2, ['Recommendation', 'Count'])
                cursor.execute("SELECT recommendation, COUNT(*) as c FROM candidate_scores GROUP BY recommendation")
                for row in cursor.fetchall():
                    ws2.append([row['recommendation'], row['c']])
                self._auto_adjust_width(ws2)

                # 03_Recruitment_Summary
                ws3 = wb.create_sheet("03_Recruitment_Summary")
                self._create_header(ws3, ['Job ID', 'Job Title', 'Company', 'Recruiter', 'Department', 'Employment Type', 'Location', 'Deadline', 'Vacancies', 'Uploaded', 'AI Processed', 'Failed', 'Shortlisted', 'Waitlisted', 'Rejected', 'Avg AI Score', 'Avg Similarity'])
                cursor.execute('''
                    SELECT j.*, r.full_name as recruiter_name,
                        (SELECT COUNT(*) FROM resumes res WHERE res.job_id = j.job_id) as uploaded,
                        (SELECT COUNT(*) FROM resumes res WHERE res.job_id = j.job_id AND processing_status = 'Completed') as processed,
                        (SELECT COUNT(*) FROM resumes res WHERE res.job_id = j.job_id AND processing_status = 'Failed') as failed,
                        (SELECT COUNT(*) FROM candidate_selection cs WHERE cs.job_id = j.job_id AND cs.status = 'Shortlisted') as short,
                        (SELECT COUNT(*) FROM candidate_selection cs WHERE cs.job_id = j.job_id AND cs.status = 'Waitlisted') as wait,
                        (SELECT COUNT(*) FROM candidate_selection cs WHERE cs.job_id = j.job_id AND cs.status = 'Rejected') as rej,
                        (SELECT COALESCE(AVG(overall_score), 0) FROM candidate_scores s WHERE s.job_id = j.job_id) as avg_score,
                        (SELECT COALESCE(AVG(semantic_score), 0) FROM candidate_scores s WHERE s.job_id = j.job_id) as avg_sem
                    FROM job_descriptions j
                    LEFT JOIN recruiters r ON j.recruiter_id = r.recruiter_id
                ''')
                for row in cursor.fetchall():
                    ws3.append([row['job_id'], row['job_title'], row['company'], row['recruiter_name'], row['department'], row['employment_type'], row['location'], str(row['application_deadline']), row['vacancies'], row['uploaded'], row['processed'], row['failed'], row['short'], row['wait'], row['rej'], int(row['avg_score']), int(row['avg_sem'])])
                self._auto_adjust_width(ws3)

                # 04_Recruitment_Details
                ws4 = wb.create_sheet("04_Recruitment_Details")
                self._create_header(ws4, ['Job ID', 'Job Title', 'Candidate ID', 'Candidate Name', 'Recruiter', 'Resume ID', 'Resume File', 'AI Score', 'Semantic Score', 'Recommendation', 'Decision', 'Decision Date'])
                cursor.execute('''
                    SELECT res.job_id, j.job_title, c.candidate_id, c.full_name, r.full_name as recruiter, res.resume_id, res.file_name, cs.overall_score, cs.semantic_score, cs.recommendation, sel.status, sel.decision_date
                    FROM resumes res
                    LEFT JOIN job_descriptions j ON res.job_id = j.job_id
                    LEFT JOIN candidates c ON res.candidate_id = c.candidate_id
                    LEFT JOIN recruiters r ON res.recruiter_id = r.recruiter_id
                    LEFT JOIN candidate_scores cs ON res.resume_id = cs.resume_id
                    LEFT JOIN candidate_selection sel ON res.resume_id = sel.resume_id
                ''')
                for row in cursor.fetchall():
                    ws4.append([row['job_id'], row['job_title'], row['candidate_id'], row['full_name'], row['recruiter'], row['resume_id'], row['file_name'], row['overall_score'], row['semantic_score'], row['recommendation'], row['status'], str(row['decision_date']) if row['decision_date'] else ''])
                self._auto_adjust_width(ws4)

                # 05_Recruiter_Summary
                ws5 = wb.create_sheet("05_Recruiter_Summary")
                self._create_header(ws5, ['Recruiter ID', 'Recruiter Name', 'Company', 'Email', 'Jobs Created', 'Uploaded', 'AI Processed', 'Failed', 'Shortlisted', 'Waitlisted', 'Rejected', 'Avg AI Score'])
                cursor.execute('''
                    SELECT r.recruiter_id, r.full_name, r.company, r.email,
                        (SELECT COUNT(*) FROM job_descriptions j WHERE j.recruiter_id = r.recruiter_id) as jobs,
                        (SELECT COUNT(*) FROM resumes res WHERE res.recruiter_id = r.recruiter_id) as uploaded,
                        (SELECT COUNT(*) FROM resumes res WHERE res.recruiter_id = r.recruiter_id AND processing_status = 'Completed') as processed,
                        (SELECT COUNT(*) FROM resumes res WHERE res.recruiter_id = r.recruiter_id AND processing_status = 'Failed') as failed,
                        (SELECT COUNT(*) FROM candidate_selection cs WHERE cs.recruiter_id = r.recruiter_id AND cs.status = 'Shortlisted') as short,
                        (SELECT COUNT(*) FROM candidate_selection cs WHERE cs.recruiter_id = r.recruiter_id AND cs.status = 'Waitlisted') as wait,
                        (SELECT COUNT(*) FROM candidate_selection cs WHERE cs.recruiter_id = r.recruiter_id AND cs.status = 'Rejected') as rej,
                        (SELECT COALESCE(AVG(cs.overall_score), 0) FROM candidate_scores cs JOIN resumes res ON cs.resume_id = res.resume_id WHERE res.recruiter_id = r.recruiter_id) as avg_score
                    FROM recruiters r
                ''')
                for row in cursor.fetchall():
                    ws5.append([row['recruiter_id'], row['full_name'], row['company'], row['email'], row['jobs'], row['uploaded'], row['processed'], row['failed'], row['short'], row['wait'], row['rej'], int(row['avg_score'])])
                self._auto_adjust_width(ws5)

                # 06_Recruiter_Details
                ws6 = wb.create_sheet("06_Recruiter_Details")
                self._create_header(ws6, ['Recruiter', 'Job ID', 'Job Title', 'Uploaded', 'Processed', 'Failed', 'Shortlisted', 'Waitlisted', 'Rejected'])
                cursor.execute('''
                    SELECT r.full_name as recruiter, j.job_id, j.job_title,
                        (SELECT COUNT(*) FROM resumes res WHERE res.job_id = j.job_id AND res.recruiter_id = r.recruiter_id) as uploaded,
                        (SELECT COUNT(*) FROM resumes res WHERE res.job_id = j.job_id AND res.recruiter_id = r.recruiter_id AND processing_status = 'Completed') as processed,
                        (SELECT COUNT(*) FROM resumes res WHERE res.job_id = j.job_id AND res.recruiter_id = r.recruiter_id AND processing_status = 'Failed') as failed,
                        (SELECT COUNT(*) FROM candidate_selection cs WHERE cs.job_id = j.job_id AND cs.recruiter_id = r.recruiter_id AND cs.status = 'Shortlisted') as short,
                        (SELECT COUNT(*) FROM candidate_selection cs WHERE cs.job_id = j.job_id AND cs.recruiter_id = r.recruiter_id AND cs.status = 'Waitlisted') as wait,
                        (SELECT COUNT(*) FROM candidate_selection cs WHERE cs.job_id = j.job_id AND cs.recruiter_id = r.recruiter_id AND cs.status = 'Rejected') as rej
                    FROM recruiters r
                    JOIN job_descriptions j ON r.recruiter_id = j.recruiter_id
                ''')
                for row in cursor.fetchall():
                    ws6.append([row['recruiter'], row['job_id'], row['job_title'], row['uploaded'], row['processed'], row['failed'], row['short'], row['wait'], row['rej']])
                self._auto_adjust_width(ws6)

                # 07_AI_Summary
                ws7 = wb.create_sheet("07_AI_Summary")
                self._create_header(ws7, ['Metric', 'Value'])
                cursor.execute("SELECT COALESCE(AVG(semantic_score), 0) as s FROM candidate_scores")
                sem = int(cursor.fetchone()['s'])
                ws7.append(['Total Resumes Processed', ai_processed])
                ws7.append(['Avg AI Score', f"{avg_score}%"])
                ws7.append(['Avg Semantic Similarity', f"{sem}%"])
                self._auto_adjust_width(ws7)

                # 08_AI_Job_Details
                ws8 = wb.create_sheet("08_AI_Job_Details")
                self._create_header(ws8, ['Job ID', 'Job Title', 'Uploaded', 'Processed', 'Failed', 'Highest Score', 'Lowest Score', 'Average AI Score'])
                cursor.execute('''
                    SELECT j.job_id, j.job_title,
                        (SELECT COUNT(*) FROM resumes res WHERE res.job_id = j.job_id) as uploaded,
                        (SELECT COUNT(*) FROM resumes res WHERE res.job_id = j.job_id AND processing_status = 'Completed') as processed,
                        (SELECT COUNT(*) FROM resumes res WHERE res.job_id = j.job_id AND processing_status = 'Failed') as failed,
                        (SELECT COALESCE(MAX(overall_score), 0) FROM candidate_scores cs WHERE cs.job_id = j.job_id) as highest,
                        (SELECT COALESCE(MIN(overall_score), 0) FROM candidate_scores cs WHERE cs.job_id = j.job_id) as lowest,
                        (SELECT COALESCE(AVG(overall_score), 0) FROM candidate_scores cs WHERE cs.job_id = j.job_id) as avg_score
                    FROM job_descriptions j
                ''')
                for row in cursor.fetchall():
                    ws8.append([row['job_id'], row['job_title'], row['uploaded'], row['processed'], row['failed'], row['highest'], row['lowest'], int(row['avg_score'])])
                self._auto_adjust_width(ws8)

                # 09_Candidate_Evaluation
                ws9 = wb.create_sheet("09_Candidate_Evaluation")
                self._create_header(ws9, ['Candidate ID', 'Candidate Name', 'Job ID', 'Job Title', 'Recruiter', 'Resume ID', 'Skills Score', 'Education Score', 'Experience Score', 'Semantic Score', 'Final Score', 'Recommendation', 'Evaluation Date'])
                cursor.execute('''
                    SELECT c.candidate_id, c.full_name, res.job_id, j.job_title, r.full_name as recruiter, res.resume_id, cs.skills_score, cs.education_score, cs.experience_score, cs.semantic_score, cs.overall_score, cs.recommendation, cs.created_at
                    FROM candidate_scores cs
                    JOIN resumes res ON cs.resume_id = res.resume_id
                    LEFT JOIN candidates c ON res.candidate_id = c.candidate_id
                    LEFT JOIN job_descriptions j ON res.job_id = j.job_id
                    LEFT JOIN recruiters r ON res.recruiter_id = r.recruiter_id
                ''')
                for row in cursor.fetchall():
                    ws9.append([row['candidate_id'], row['full_name'], row['job_id'], row['job_title'], row['recruiter'], row['resume_id'], row['skills_score'], row['education_score'], row['experience_score'], row['semantic_score'], row['overall_score'], row['recommendation'], str(row['created_at'])])
                self._auto_adjust_width(ws9)

                # 10_Resume_Processing
                ws10 = wb.create_sheet("10_Resume_Processing")
                self._create_header(ws10, ['Resume ID', 'Candidate ID', 'Job ID', 'Recruiter', 'File Name', 'File Type', 'File Size', 'Upload Date', 'Status'])
                cursor.execute('''
                    SELECT res.resume_id, res.candidate_id, res.job_id, r.full_name as recruiter, res.file_name, res.file_type, res.file_size, res.upload_date, res.processing_status
                    FROM resumes res
                    LEFT JOIN recruiters r ON res.recruiter_id = r.recruiter_id
                ''')
                for row in cursor.fetchall():
                    ws10.append([row['resume_id'], row['candidate_id'], row['job_id'], row['recruiter'], row['file_name'], row['file_type'], row['file_size'], str(row['upload_date']), row['processing_status']])
                self._auto_adjust_width(ws10)

                # 11_Semantic_Scores
                ws11 = wb.create_sheet("11_Semantic_Scores")
                self._create_header(ws11, ['Semantic ID', 'Resume ID', 'Candidate ID', 'Job Title', 'Semantic Similarity', 'Matched Keywords'])
                cursor.execute('''
                    SELECT s.semantic_id, s.resume_id, res.candidate_id, j.job_title, s.semantic_score, s.matched_keywords
                    FROM semantic_scores s
                    JOIN resumes res ON s.resume_id = res.resume_id
                    JOIN job_descriptions j ON res.job_id = j.job_id
                ''')
                for row in cursor.fetchall():
                    ws11.append([row.get('semantic_id', ''), row['resume_id'], row['candidate_id'], row['job_title'], row['semantic_score'], row['matched_keywords']])
                self._auto_adjust_width(ws11)

                # 12_LinkedIn_Analysis
                ws12 = wb.create_sheet("12_LinkedIn_Analysis")
                self._create_header(ws12, ['Resume ID', 'Candidate ID', 'LinkedIn URL', 'Profile Completeness', 'LinkedIn Score'])
                cursor.execute('''
                    SELECT l.resume_id, res.candidate_id, l.linkedin_url, l.profile_completeness, l.linkedin_score
                    FROM linkedin_analysis l
                    JOIN resumes res ON l.resume_id = res.resume_id
                ''')
                for row in cursor.fetchall():
                    ws12.append([row['resume_id'], row['candidate_id'], row['linkedin_url'], row['profile_completeness'], row['linkedin_score']])
                self._auto_adjust_width(ws12)

                # 13_Portfolio_Analysis
                ws13 = wb.create_sheet("13_Portfolio_Analysis")
                self._create_header(ws13, ['Resume ID', 'Candidate ID', 'Portfolio URL', 'Portfolio Score', 'Analysis Summary'])
                cursor.execute('''
                    SELECT p.resume_id, res.candidate_id, p.portfolio_url, p.portfolio_score, p.analysis_summary
                    FROM portfolio_analysis p
                    JOIN resumes res ON p.resume_id = res.resume_id
                ''')
                for row in cursor.fetchall():
                    ws13.append([row['resume_id'], row['candidate_id'], row['portfolio_url'], row['portfolio_score'], row['analysis_summary']])
                self._auto_adjust_width(ws13)

                # 14_Candidate_Decisions
                ws14 = wb.create_sheet("14_Candidate_Decisions")
                self._create_header(ws14, ['Selection ID', 'Candidate ID', 'Candidate Name', 'Job Title', 'Recruiter', 'Decision', 'Decision Date'])
                cursor.execute('''
                    SELECT sel.selection_id, c.candidate_id, c.full_name, j.job_title, r.full_name as recruiter, sel.status, sel.decision_date
                    FROM candidate_selection sel
                    JOIN candidates c ON sel.candidate_id = c.candidate_id
                    JOIN job_descriptions j ON sel.job_id = j.job_id
                    JOIN recruiters r ON sel.recruiter_id = r.recruiter_id
                ''')
                for row in cursor.fetchall():
                    ws14.append([row['selection_id'], row['candidate_id'], row['full_name'], row['job_title'], row['recruiter'], row['status'], str(row['decision_date']) if row['decision_date'] else ''])
                self._auto_adjust_width(ws14)

                # 15_Activity_Logs
                ws15 = wb.create_sheet("15_Activity_Logs")
                self._create_header(ws15, ['Log ID', 'User ID', 'Action', 'Module', 'Description', 'Timestamp'])
                cursor.execute('''
                    SELECT log_id, recruiter_id, admin_id, action, module, action_time
                    FROM activity_logs
                    ORDER BY action_time DESC
                ''')
                for row in cursor.fetchall():
                    user_id = row['recruiter_id'] if row['recruiter_id'] else (row['admin_id'] if row['admin_id'] else 'System')
                    ws15.append([row['log_id'], user_id, row['action'], row['module'], row['action'], str(row['action_time'])])
                self._auto_adjust_width(ws15)

        except Exception as e:
            print(f"Error generating excel: {e}")
        finally:
            conn.close()
            
        memory_file = io.BytesIO()
        wb.save(memory_file)
        memory_file.seek(0)
        return memory_file

reports_export_service = ReportsExportService()
